"""Preenche cópias de relatório e slides SOMENTE de experimento Satimage real.

Uso: python gerar_documentos_com_resultados.py --results results/imbalance_YYYYmmdd_HHMMSS --budget 20
Não altera documentos modelos e recusa smoke/quick. Checa quatro braços completos
com sementes pareadas antes de produzir qualquer saída.
"""
from __future__ import annotations
import argparse
import json
from pathlib import Path
import pandas as pd
from docx import Document
from pptx import Presentation

ROOT = Path(__file__).resolve().parent
ARMS = [('accuracy','none'), ('f1_macro','none'),
        ('accuracy','balanced'), ('f1_macro','balanced')]
ARM_NAMES = ['A','B','C','D']
METHOD_NAMES = {'pso':'PSO', 'adaptive_pso':'Adaptive PSO'}


def validate(run_dir: Path, budget: int):
    manifest=json.loads((run_dir/'experiment_manifest.json').read_text(encoding='utf-8'))
    description=json.loads((run_dir/'dataset_report.json').read_text(encoding='utf-8'))
    if manifest.get('mode')!='imbalance' or manifest.get('smoke') or manifest.get('quick'):
        raise ValueError('Recusado: exige --mode imbalance, base REAL e execução sem --smoke/--quick.')
    if ((description.get('samples'),description.get('attributes'),description.get('n_classes')) != (6435,36,6)
        or not description.get('sha256')):
        raise ValueError('Recusado: integridade/metadados da Satimage 6435×36/6 classes não confirmados.')
    df=pd.read_csv(run_dir/'runs.csv')
    sub=df[(df['budget']==budget) & (df['status']=='ok')].copy()
    if len(sub)==0:raise ValueError(f'Não há execuções concluídas para orçamento {budget}.')
    refs={}
    for method in METHOD_NAMES:
        m=sub[sub['method']==method]
        arms={}
        for (fitness,weight),label in zip(ARMS,ARM_NAMES):
            part=m[(m['fitness_metric']==fitness)&(m['class_weight']==weight)]
            if part.empty or part['seed'].duplicated().any():
                raise ValueError(f'Braço incompleto/duplicado: {method} {label} orçamento={budget}')
            arms[label]=part.set_index('seed')
        matching=set(arms['A'].index)
        if not matching or any(set(z.index)!=matching for z in arms.values()):
            raise ValueError(f'Sementes dos braços A–D não coincidem no método {method}.')
        refs[method]=arms
    return refs,description


def mean_std(ser):
    mean=float(ser.mean());std=float(ser.std(ddof=1)) if len(ser)>1 else float('nan')
    return f'{mean:.4f} ± {std:.4f}' if pd.notna(std) else f'{mean:.4f} (n=1)'


def put_text(sh, val):
    tf=sh.text_frame
    if tf.paragraphs and tf.paragraphs[0].runs:
        tf.paragraphs[0].runs[0].text=val
        for p in tf.paragraphs:
            for r in (p.runs[1:] if p is tf.paragraphs[0] else p.runs):r.text=''
    else:sh.text=val


def write_docs(refs, description, run_dir, budget, method):
    output=run_dir/'documentos_preenchidos';output.mkdir(exist_ok=True)
    doc=Document(ROOT/'docs'/'RELATORIO_MINIPROJETO_EDITAVEL.docx')
    target=None
    for table in doc.tables:
        if table.cell(0,0).text=='Método' and table.cell(0,1).text=='Braço':
            target=table;break
    if target is None:raise RuntimeError('Tabela de quatro braços ausente do relatório modelo.')
    for i,meth in enumerate(METHOD_NAMES):
        for j,label in enumerate(ARM_NAMES):
            row=target.rows[1+i*4+j].cells
            vals=refs[meth][label]
            for k,col in enumerate(['test_f1_macro','test_accuracy','test_min_class_recall','elapsed_total_s'],start=2):
                row[k].text=mean_std(vals[col])
    # Tabela geral: apenas os PSOs possuem resultados no modo imbalance.
    # Para TPE/Random Search, não extrapolar valores de métodos não executados.
    base_table = next((table for table in doc.tables
        if table.cell(0, 0).text == 'Método'
        and 'Acurácia média' in table.cell(0, 1).text), None)
    if base_table is not None:
        for cells in (r.cells for r in base_table.rows[1:]):
            name = cells[0].text.strip()
            key = {'PSO':'pso', 'PSO adaptativo':'adaptive_pso'}.get(name)
            if key:
                arm = refs[key]['A']
                for col_idx, field in [(1,'test_accuracy'),(2,'test_f1_macro'),
                                       (3,'elapsed_total_s')]:
                    cells[col_idx].text = mean_std(arm[field])
            else:
                for cell in cells[1:]:
                    cell.text = 'Não avaliado neste modo'
    for p in doc.paragraphs:
        if p.text.startswith('Esta tabela é um modelo;'):
            p.text=(f'Tabela preenchida automaticamente de {run_dir.name}, budget {budget}, '
                    f'dataset sha256={description["sha256"][:16]}…; '
                    f'mesma(s) semente(s) em A–D. Dados experimentais; conferir outros budgets separadamente.')
        if p.text.startswith('[INSERIR APÓS EXECUTAR SATIMAGE REAL]'):
            p.text=(f'Tabela 1 — Satimage real, orçamento {budget}; resultados PSO e '
                    'PSO adaptativo no braço A (accuracy/none). TPE/Random Search não '
                    'avaliados neste modo; executar extensão separada para compará-los.')
        if p.text.startswith('[INSERIR RESULTADOS SATIMAGE]'):
            p.text=(f'Tabela fatorial A–D preenchida com resultados reais, orçamento {budget}. '
                    'Comparar cada método sob as mesmas sementes; não interpretar o melhor '
                    'cenário do teste como seleção independente de modelo.')
        if p.text.startswith('Conclusão provisória:'):
            mean_diff=(refs['adaptive_pso']['D']['test_f1_macro']
                       - refs['adaptive_pso']['A']['test_f1_macro']).mean()
            p.text=(f'Conclusão experimental delimitada: na Satimage, no orçamento {budget}, '
                    f'a diferença média D−A do PSO adaptativo em F1-macro foi '
                    f'{mean_diff:+.4f} nas sementes registradas. Esta observação não '
                    f'comprova ganho geral; verificar as classes, a variância e '
                    f'outros orçamentos antes de qualquer conclusão abrangente.')
        if p.text.startswith('ESTADO DA IMPLEMENTAÇÃO:'):
            p.text=(f'RESULTADOS INCLUÍDOS: experimento Satimage REAL; orçamento {budget}. '
                    f'Base validada: 6435 amostras, 36 atributos, 6 classes; hash '
                    f'{description["sha256"][:16]}….')
    output_doc=output/f'RELATORIO_SATIMAGE_REAL_budget{budget}.docx'
    doc.save(output_doc)
    ppt=Presentation(ROOT/'docs'/'APRESENTACAO_6_INTEGRANTES_EDITAVEL.pptx')
    slide=ppt.slides[28]
    put_text(slide.shapes[1],f'INTEGRANTE 6  •  SATIMAGE REAL  •  {METHOD_NAMES[method].upper()}')
    put_text(slide.shapes[2],f'F1-macro: comparação A–D ({METHOD_NAMES[method]}, {budget} aval.)')
    put_text(slide.shapes[7],f'EXPERIMENTO REAL | {len(refs[method]["A"])} sementes | conjuntos de teste idênticos')
    for label,idx in zip(['B','C','D'],[10,13,16]):
        delta=refs[method][label]['test_f1_macro']-refs[method]['A']['test_f1_macro']
        put_text(slide.shapes[idx],f'{delta.mean():+.4f} (média; n={len(delta)})')
    # Interação fatorial: (D-B-C+A), calculada em sementes comuns.
    effect=(refs[method]['D']['test_f1_macro'] - refs[method]['B']['test_f1_macro']
           - refs[method]['C']['test_f1_macro'] + refs[method]['A']['test_f1_macro'])
    put_text(slide.shapes[18],
             f'Interação D−B−C+A = {effect.mean():+.4f} | Fonte: runs.csv')
    slide2=ppt.slides[29]
    rows=refs[method]
    summaries=[]
    for label in ['A','D']:
        x=rows[label]
        summaries.append((label,float(x['test_accuracy'].mean()),float(x['test_min_class_recall'].mean())))
    put_text(slide2.shapes[7],f'A: acurácia={summaries[0][1]:.4f}, recall mínimo={summaries[0][2]:.4f}')
    put_text(slide2.shapes[9],f'D: acurácia={summaries[1][1]:.4f}, recall mínimo={summaries[1][2]:.4f}')
    put_text(slide2.shapes[11],'Atenção: comparar também as seis classes em per_class/*.csv.')
    put_text(slide2.shapes[13],f'Orçamento {budget} | {len(rows["A"])} sementes | fonte: Satimage REAL')
    output_ppt=output/f'APRESENTACAO_SATIMAGE_REAL_budget{budget}.pptx'
    ppt.save(output_ppt)
    meta={'dataset_sha256':description['sha256'],'source_directory':str(run_dir.resolve()),
          'budget':budget,'method_ppt':method,'n_paired_seeds':len(refs[method]['A']),
          'output_docx':str(output_doc),'output_pptx':str(output_ppt)}
    (output/'proveniencia.json').write_text(json.dumps(meta,indent=2,ensure_ascii=False),encoding='utf-8')
    print('Relatório:',output_doc)
    print('Slides:',output_ppt)
    print('Proveniência:',output/'proveniencia.json')


def main():
    ap=argparse.ArgumentParser()
    ap.add_argument('--results',required=True,help='Pasta de saída de execução --mode imbalance REAL')
    ap.add_argument('--budget',required=True,type=int,help='Orçamento a exibir, como 20 ou 50')
    ap.add_argument('--method',choices=METHOD_NAMES,default='adaptive_pso',help='Método exibido no slide de diferenças')
    args=ap.parse_args()
    path=Path(args.results)
    refs,info=validate(path,args.budget)
    write_docs(refs,info,path,args.budget,args.method)
if __name__=='__main__': main()
