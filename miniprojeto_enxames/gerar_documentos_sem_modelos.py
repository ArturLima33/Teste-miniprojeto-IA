import json
from pathlib import Path

import pandas as pd
from docx import Document
from pptx import Presentation
from pptx.util import Inches, Pt

RUN = Path("results/imbalance_20261028_003143")
OUT = RUN / "documentos_gerados"
OUT.mkdir(parents=True, exist_ok=True)

df = pd.read_csv(RUN / "aggregated.csv")
info = json.loads((RUN / "dataset_report.json").read_text(encoding="utf-8"))

required = {
    "test_accuracy_mean", "test_f1_macro_mean",
    "test_min_class_recall_mean", "budget", "method",
    "fitness_metric", "class_weight"
}
missing = required - set(df.columns)
if missing:
    raise ValueError(f"Colunas em falta: {sorted(missing)}")

def label(row):
    return (
        f'{row["method"]} | orçamento {row["budget"]} | '
        f'objetivo {row["fitness_metric"]} | pesos {row["class_weight"]}'
    )

cols = [
    "budget", "method", "fitness_metric", "class_weight",
    "test_accuracy_mean", "test_f1_macro_mean",
    "test_balanced_accuracy_mean", "test_min_class_recall_mean"
]

best_f1 = df.loc[df["test_f1_macro_mean"].idxmax()]
best_acc = df.loc[df["test_accuracy_mean"].idxmax()]
best_recall = df.loc[df["test_min_class_recall_mean"].idxmax()]

# ---------- Relatório Word ----------
doc = Document()
doc.add_heading("Miniprojeto de Inteligência Artificial", 0)
doc.add_heading("Classificação de imagens de satélite — Satimage", 1)

doc.add_paragraph(
    "Relatório gerado automaticamente a partir dos resultados experimentais "
    "guardados pelo projeto. A análise compara PSO e Adaptive PSO, "
    "dois orçamentos de avaliação, duas funções objetivo e duas opções "
    "de ponderação das classes."
)

doc.add_heading("1. Dados utilizados", 1)
doc.add_paragraph(f'Amostras: {info["samples"]}')
doc.add_paragraph(f'Atributos: {info["attributes"]}')
doc.add_paragraph(f'Número de classes: {info["n_classes"]}')
doc.add_paragraph(f'Rácio de desbalanceamento: {info["imbalance_ratio"]:.4f}')
doc.add_paragraph(f'Ficheiro: {info["source"]}')
doc.add_paragraph(f'SHA-256: {info["sha256"]}')
doc.add_paragraph("Contagem por classe:")
for cls, count in info["class_counts"].items():
    doc.add_paragraph(f"Classe {cls}: {count}", style="List Bullet")

doc.add_heading("2. Melhores configurações observadas", 1)
for title, row, metric in [
    ("Maior F1-macro médio", best_f1, "test_f1_macro_mean"),
    ("Maior acurácia média", best_acc, "test_accuracy_mean"),
    ("Maior recall mínimo médio", best_recall, "test_min_class_recall_mean"),
]:
    doc.add_heading(title, 2)
    doc.add_paragraph(label(row))
    doc.add_paragraph(f'{metric}: {row[metric]:.4f}')

doc.add_heading("3. Comparação das configurações", 1)
table = doc.add_table(rows=1, cols=len(cols))
table.style = "Table Grid"
for cell, col in zip(table.rows[0].cells, cols):
    cell.text = col
for _, row in df.sort_values("test_f1_macro_mean", ascending=False).iterrows():
    cells = table.add_row().cells
    for cell, col in zip(cells, cols):
        value = row[col]
        cell.text = f"{value:.4f}" if isinstance(value, float) else str(value)

doc.add_heading("4. Interpretação e limitações", 1)
doc.add_paragraph(
    "A configuração com maior F1-macro não tem necessariamente o maior "
    "recall para a classe mais difícil. A ponderação balanced pode favorecer "
    "o recall mínimo e reduzir a acurácia global. A escolha deve considerar "
    "o custo dos erros por classe."
)
doc.add_paragraph(
    "Os resultados são médias agregadas de cinco sementes por configuração. "
    "As sementes partilham a mesma divisão externa dos dados; portanto, "
    "não representam cinco amostras independentes do conjunto de dados. "
    "As comparações no teste são exploratórias e não demonstram, por si só, "
    "superioridade estatística nem generalização para outras divisões."
)
doc.add_paragraph(
    "Segundo a interpretação guardada pelo experimento, SMOTE e Tomek links "
    "não foram aplicados. A seleção de hiperparâmetros não utiliza o conjunto "
    "de teste."
)

docx_path = OUT / "RELATORIO_SATIMAGE_REAL.docx"
doc.save(docx_path)

# ---------- Apresentação PowerPoint ----------
prs = Presentation()

def add_bullet_slide(title, lines):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    tf = slide.placeholders[1].text_frame
    tf.clear()
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.size = Pt(20)
    return slide

add_bullet_slide("Classificação de imagens de satélite", [
    "Miniprojeto de Inteligência Artificial",
    "Dataset Satimage",
    "Comparação entre PSO e Adaptive PSO",
    "Experimento real — resultados agregados"
])

add_bullet_slide("Dados e metodologia", [
    f'Amostras: {info["samples"]}; atributos: {info["attributes"]}',
    f'Classes: {info["n_classes"]}; rácio de desbalanceamento: {info["imbalance_ratio"]:.3f}',
    "Orçamentos: 20 e 50 avaliações por execução",
    "Objetivos: accuracy e F1-macro",
    "Pesos de classe: none e balanced",
    "Cinco sementes por configuração"
])

add_bullet_slide("Principais resultados", [
    f'Maior F1-macro: {best_f1["test_f1_macro_mean"]:.4f}',
    label(best_f1),
    f'Maior acurácia: {best_acc["test_accuracy_mean"]:.4f}',
    label(best_acc),
    f'Maior recall mínimo: {best_recall["test_min_class_recall_mean"]:.4f}',
    label(best_recall)
])

add_bullet_slide("Conclusões e limitações", [
    "A melhor configuração depende da métrica prioritária.",
    "class_weight=balanced pode melhorar o recall da classe mais difícil.",
    "As métricas são médias de cinco sementes na mesma divisão externa.",
    "Comparações no teste são exploratórias.",
    "Uma única divisão não demonstra superioridade geral."
])

pptx_path = OUT / "APRESENTACAO_SATIMAGE_REAL.pptx"
prs.save(pptx_path)

print("Documentos criados:")
print(docx_path.resolve())
print(pptx_path.resolve())
print("Fonte dos dados:", RUN.resolve())
